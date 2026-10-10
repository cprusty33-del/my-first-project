"""Read one SAP table (view only) and save it as an Excel file for TeamMate Analytics.

Uses the standard SAP function RFC_READ_TABLE through PyRFC. Nothing is written
to SAP. The password is asked for on screen each time and is never saved.
No AI service is used: the program talks only to your SAP server.
"""

import argparse
import configparser
import datetime as dt
import getpass
import os
import re
import sys

ROW_WIDTH = 512          # RFC_READ_TABLE returns at most 512 characters per row
OPTION_WIDTH = 72        # each WHERE line sent to SAP is at most 72 characters
EXCEL_MAX_ROWS = 1_048_575
NUMBER_TYPES = {"P", "I", "F", "b", "s", "8"}

HERE = os.path.dirname(os.path.abspath(sys.argv[0] if getattr(sys, "frozen", False) else __file__))


def split_where(where):
    """Break a WHERE condition into lines of at most 72 characters, at spaces."""
    lines, current = [], ""
    for word in where.split():
        if len(word) > OPTION_WIDTH:
            raise ValueError(f"A single word in the filter is longer than {OPTION_WIDTH} characters: {word}")
        candidate = f"{current} {word}" if current else word
        if len(candidate) > OPTION_WIDTH:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return [{"TEXT": line} for line in lines]


def read_field_list(conn, table):
    """Ask SAP for the table's fields (name, length, type, description) without reading data."""
    result = conn.call("RFC_READ_TABLE", QUERY_TABLE=table, NO_DATA="X")
    return [
        {
            "name": f["FIELDNAME"].strip(),
            "length": int(f["LENGTH"]),
            "type": f["TYPE"].strip(),
            "text": f["FIELDTEXT"].strip(),
        }
        for f in result["FIELDS"]
    ]


def read_key_fields(conn, table):
    """Return the table's key field names, or None if this user may not look them up."""
    try:
        result = conn.call("DDIF_FIELDINFO_GET", TABNAME=table)
    except Exception:
        return None
    return [f["FIELDNAME"].strip() for f in result["DFIES_TAB"] if f["KEYFLAG"] == "X" and f["FIELDNAME"].strip() != ".INCLUDE"]


def plan_groups(fields, key_names):
    """Split fields into groups that each fit in 512 characters.

    When more than one group is needed, every group also carries the key
    fields so the pieces can be joined back into whole rows safely.
    """
    if sum(f["length"] for f in fields) <= ROW_WIDTH:
        return [fields]
    if not key_names:
        raise ValueError(
            "The chosen fields are wider than SAP allows in one read (512 characters), "
            "and the table's key fields could not be looked up. Choose fewer fields."
        )
    by_name = {f["name"]: f for f in fields}
    keys = [by_name[k] for k in key_names if k in by_name]
    if len(keys) != len(key_names):
        raise ValueError("When many fields are chosen, all key fields must be included: " + ", ".join(key_names))
    key_width = sum(f["length"] for f in keys)
    others = [f for f in fields if f["name"] not in key_names]
    groups, current, width = [], [], key_width
    for f in others:
        if key_width + f["length"] > ROW_WIDTH:
            raise ValueError(f"Field {f['name']} is too long to read together with the key fields.")
        if width + f["length"] > ROW_WIDTH:
            groups.append(keys + current)
            current, width = [], key_width
        current.append(f)
        width += f["length"]
    if current:
        groups.append(keys + current)
    return groups


def read_group(conn, table, group, options, max_rows):
    """Read one group of fields and return a list of {field: text} rows."""
    result = conn.call(
        "RFC_READ_TABLE",
        QUERY_TABLE=table,
        DELIMITER="",
        ROWCOUNT=max_rows,
        OPTIONS=options,
        FIELDS=[{"FIELDNAME": f["name"]} for f in group],
    )
    layout = [(f["FIELDNAME"].strip(), int(f["OFFSET"]), int(f["LENGTH"])) for f in result["FIELDS"]]
    rows = []
    for line in result["DATA"]:
        wa = line["WA"]
        rows.append({name: wa[offset:offset + length].strip() for name, offset, length in layout})
    return rows


def convert(value, sap_type):
    """Turn SAP text into a number or date where that is safe; otherwise keep the text."""
    if sap_type in NUMBER_TYPES:
        text = value.replace(" ", "")
        if text.endswith("-"):
            text = "-" + text[:-1]
        try:
            number = float(text) if text else None
        except ValueError:
            return value
        if number is not None and number.is_integer() and "." not in text and "E" not in text.upper():
            return int(number)
        return number
    if sap_type == "D":
        if re.fullmatch(r"\d{8}", value) and value != "00000000":
            try:
                return dt.date(int(value[:4]), int(value[4:6]), int(value[6:]))
            except ValueError:
                return value
        return None if value in ("", "00000000") else value
    return value


def extract(conn, table, field_names, where, max_rows):
    """Read the table and return (fields, rows) with values converted."""
    all_fields = read_field_list(conn, table)
    if field_names:
        by_name = {f["name"]: f for f in all_fields}
        missing = [n for n in field_names if n not in by_name]
        if missing:
            raise ValueError(f"These fields are not in table {table}: {', '.join(missing)}")
        fields = [by_name[n] for n in field_names]
    else:
        fields = all_fields

    key_names = None
    if sum(f["length"] for f in fields) > ROW_WIDTH:
        key_names = read_key_fields(conn, table)
    groups = plan_groups(fields, key_names)
    options = split_where(where) if where else []

    pieces = [read_group(conn, table, g, options, max_rows) for g in groups]
    if len(pieces) == 1:
        merged = pieces[0]
    else:
        merged = pieces[0]
        index_of = {}
        for i, row in enumerate(merged):
            index_of[tuple(row[k] for k in key_names)] = i
        for piece in pieces[1:]:
            if len(piece) != len(merged):
                raise RuntimeError("The table changed while it was being read. Please run again.")
            for row in piece:
                key = tuple(row[k] for k in key_names)
                if key not in index_of:
                    raise RuntimeError("The table changed while it was being read. Please run again.")
                merged[index_of[key]].update(row)

    types = {f["name"]: f["type"] for f in fields}
    rows = [[convert(r[f["name"]], types[f["name"]]) for f in fields] for r in merged]
    return fields, rows


def save_excel(path, table, fields, rows, where, user, system_label):
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    data = wb.create_sheet(table[:31])
    data.append([f["name"] for f in fields])
    for row in rows:
        data.append(row)

    info = wb.create_sheet("Fields")
    info.append(["Field", "Description", "SAP type", "Length"])
    for f in fields:
        info.append([f["name"], f["text"], f["type"], f["length"]])

    log = wb.create_sheet("Extract details")
    log.append(["Item", "Value"])
    log.append(["SAP table", table])
    log.append(["SAP system", system_label])
    log.append(["SAP user", user])
    log.append(["Filter (WHERE)", where or "(none)"])
    log.append(["Rows extracted", len(rows)])
    log.append(["Extracted on", dt.datetime.now().strftime("%d-%m-%Y %H:%M:%S")])
    log.append(["Method", "RFC_READ_TABLE (read only)"])
    wb.save(path)


def ask(prompt, default=""):
    shown = f"{prompt} [{default}]: " if default else f"{prompt}: "
    answer = input(shown).strip()
    return answer or default


def main():
    parser = argparse.ArgumentParser(description="Copy one SAP table into an Excel file (read only).")
    parser.add_argument("--config", default=os.path.join(HERE, "sap_config.ini"))
    parser.add_argument("--table", help="SAP table name, e.g. BKPF")
    parser.add_argument("--fields", help="Comma-separated field names; leave out for all fields")
    parser.add_argument("--where", help="Filter, e.g. BUKRS = '1000' AND GJAHR = '2025'")
    parser.add_argument("--max-rows", type=int, help="Largest number of rows to read")
    parser.add_argument("--out", help="Excel file to write")
    args = parser.parse_args()

    cfg = configparser.ConfigParser()
    if not cfg.read(args.config):
        sys.exit(f"Settings file not found: {args.config}")
    sap = cfg["sap"]
    defaults = cfg["extract"] if cfg.has_section("extract") else {}

    interactive = not args.table
    table = (args.table or ask("SAP table name (e.g. BKPF)")).upper()
    if not re.fullmatch(r"[A-Z0-9_/]{1,30}", table):
        sys.exit("Please enter a valid SAP table name.")
    fields_text = args.fields if args.fields is not None else (ask("Fields, separated by commas (Enter = all fields)") if interactive else "")
    field_names = [f.strip().upper() for f in fields_text.split(",") if f.strip()]
    where = args.where if args.where is not None else (ask("Filter (Enter = no filter), e.g. GJAHR = '2025'") if interactive else "")
    max_rows = args.max_rows or int(ask("Maximum rows", defaults.get("max_rows", "100000")) if interactive else defaults.get("max_rows", "100000"))
    max_rows = min(max_rows, EXCEL_MAX_ROWS)

    user = sap.get("user") or ask("SAP user ID")
    password = getpass.getpass("SAP password (not shown, not saved): ")

    params = {"user": user, "passwd": password, "client": sap["client"], "lang": sap.get("lang", "EN")}
    if sap.get("mshost"):
        params.update(mshost=sap["mshost"], sysid=sap["sysid"], group=sap.get("group", "PUBLIC"))
        system_label = f"{sap['sysid']} client {sap['client']}"
    else:
        params.update(ashost=sap["ashost"], sysnr=sap["sysnr"])
        system_label = f"{sap['ashost']} / {sap['sysnr']} client {sap['client']}"
    if sap.get("saprouter"):
        params["saprouter"] = sap["saprouter"]

    try:
        from pyrfc import Connection
    except ImportError:
        sys.exit("PyRFC is not installed. See README.md, step 2.")

    print("Connecting to SAP ...")
    conn = Connection(**params)
    try:
        print(f"Reading table {table} ...")
        fields, rows = extract(conn, table, field_names, where, max_rows)
    finally:
        conn.close()
        password = None

    out_dir = defaults.get("output_folder") or os.path.join(HERE, "output")
    os.makedirs(out_dir, exist_ok=True)
    out = args.out or os.path.join(out_dir, f"{table.replace('/', '_')}_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx")
    save_excel(out, table, fields, rows, where, user, system_label)
    print(f"Done. {len(rows)} rows saved to:\n{out}")
    if len(rows) == max_rows:
        print(f"Note: the limit of {max_rows} rows was reached; there may be more rows in SAP.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit("Cancelled.")
    except Exception as exc:  # show a plain message instead of a stack trace
        sys.exit(f"Error: {exc}")
