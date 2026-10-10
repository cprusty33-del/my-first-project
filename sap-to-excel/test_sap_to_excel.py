"""Offline test: a fake SAP connection stands in for the real server."""

import os
import tempfile
import unittest

import sap_to_excel as s

TABLE = [  # (name, length, type, description)
    ("BUKRS", 4, "C", "Company code"),
    ("BELNR", 10, "C", "Document number"),
    ("GJAHR", 4, "N", "Fiscal year"),
    ("BUDAT", 8, "D", "Posting date"),
    ("DMBTR", 15, "P", "Amount"),
    ("SGTXT", 400, "C", "Text"),
    ("XBLNR", 200, "C", "Reference"),
]
KEYS = ["BUKRS", "BELNR", "GJAHR"]
DATA = [
    {"BUKRS": "1000", "BELNR": "0000000001", "GJAHR": "2025", "BUDAT": "20250401", "DMBTR": "1500.50", "SGTXT": "Coal sale", "XBLNR": "INV-1"},
    {"BUKRS": "1000", "BELNR": "0000000002", "GJAHR": "2025", "BUDAT": "00000000", "DMBTR": "250.00-", "SGTXT": "Refund", "XBLNR": "CN-7"},
]


class FakeConnection:
    def __init__(self, reverse_second_read=False):
        self.calls = []
        self.reverse = reverse_second_read

    def call(self, name, **kw):
        self.calls.append((name, kw))
        if name == "DDIF_FIELDINFO_GET":
            return {"DFIES_TAB": [{"FIELDNAME": n, "KEYFLAG": "X" if n in KEYS else ""} for n, *_ in TABLE]}
        wanted = [f["FIELDNAME"] for f in kw.get("FIELDS", [])] or [n for n, *_ in TABLE]
        spec = {n: (l, t, d) for n, l, t, d in TABLE}
        fields, offset = [], 0
        for n in wanted:
            l, t, d = spec[n]
            fields.append({"FIELDNAME": n, "OFFSET": f"{offset:06d}", "LENGTH": f"{l:06d}", "TYPE": t, "FIELDTEXT": d})
            offset += l
        assert offset <= 512 or kw.get("NO_DATA") == "X", "row wider than 512"
        if kw.get("NO_DATA") == "X":
            return {"FIELDS": fields, "DATA": []}
        rows = list(DATA)
        if self.reverse and len(self.calls) > 3:
            rows.reverse()
        data = [{"WA": "".join(r[n].rjust(spec[n][0]) if spec[n][1] == "P" else r[n].ljust(spec[n][0]) for n in wanted)} for r in rows]
        return {"FIELDS": fields, "DATA": data[: kw["ROWCOUNT"]]}


class Tests(unittest.TestCase):
    def test_wide_table_is_split_and_rejoined_by_key(self):
        conn = FakeConnection(reverse_second_read=True)
        fields, rows = s.extract(conn, "BKPF", [], "GJAHR = '2025'", 100)
        self.assertGreater(sum(1 for c in conn.calls if c[0] == "RFC_READ_TABLE"), 2)
        self.assertEqual(rows[0][:5], ["1000", "0000000001", "2025", s.dt.date(2025, 4, 1), 1500.5])
        self.assertEqual(rows[1][4], -250.0)
        self.assertIsNone(rows[1][3])
        self.assertEqual(rows[1][5:], ["Refund", "CN-7"])

    def test_chosen_fields_and_excel_output(self):
        fields, rows = s.extract(FakeConnection(), "BKPF", ["BELNR", "DMBTR"], "", 100)
        self.assertEqual(rows, [["0000000001", 1500.5], ["0000000002", -250.0]])
        path = os.path.join(tempfile.mkdtemp(), "out.xlsx")
        s.save_excel(path, "BKPF", fields, rows, "", "AUDITOR", "PRD client 100")
        from openpyxl import load_workbook
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["BKPF", "Fields", "Extract details"])
        sheet = wb["BKPF"]
        self.assertEqual([c.value for c in sheet[1]], ["Document number (BELNR)", "Amount (DMBTR)"])
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:B3")
        self.assertEqual(sheet["A2"].number_format, "@")
        self.assertEqual(sheet["B3"].value, -250.0)
        self.assertEqual(sheet["B3"].number_format, "#,##0.00###")
        self.assertEqual([c.value for c in wb["Fields"][2]], ["Document number (BELNR)", "BELNR", "Document number", "Yes", "C", 10, None])

    def test_rows_are_sorted_by_key_and_dates_formatted(self):
        global DATA
        saved, DATA = DATA, list(reversed(DATA))
        try:
            fields, rows = s.extract(FakeConnection(), "BKPF", ["BELNR", "BUDAT"], "", 100)
        finally:
            DATA = saved
        self.assertEqual([r[0] for r in rows], ["0000000001", "0000000002"])
        path = os.path.join(tempfile.mkdtemp(), "out.xlsx")
        s.save_excel(path, "BKPF", fields, rows, "", "AUDITOR", "PRD", headers="code")
        from openpyxl import load_workbook
        sheet = load_workbook(path)["BKPF"]
        self.assertEqual(sheet["B1"].value, "BUDAT")
        self.assertEqual(sheet["B2"].number_format, "DD-MM-YYYY")

    def test_where_is_split_into_short_lines(self):
        where = " AND ".join(f"FIELD{i} = 'VALUE{i}'" for i in range(10))
        lines = s.split_where(where)
        self.assertTrue(all(len(l["TEXT"]) <= 72 for l in lines))
        self.assertEqual(" ".join(l["TEXT"] for l in lines), where)

    def test_unknown_field_is_reported(self):
        with self.assertRaises(ValueError):
            s.extract(FakeConnection(), "BKPF", ["NOPE"], "", 10)


if __name__ == "__main__":
    unittest.main()
